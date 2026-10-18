#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AAEON Presentation Generator
Automates populating AAEON 9-slide template using python-pptx.
"""

import os
import sys
import json
import argparse
from typing import Dict, Any, List
import pptx
from pptx.util import Pt

DEFAULT_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "resources",
    "AAEON_Template_v2.pptx"
)
FALLBACK_GDRIVE_TEMPLATE = r"G:\我的雲端硬碟\AAEON\簡報製作\AAEON_功能簡報空白模板_Template_v2.pptx"

def get_template_path(custom_path: str = None) -> str:
    if custom_path and os.path.exists(custom_path):
        return custom_path
    if os.path.exists(DEFAULT_TEMPLATE_PATH):
        return DEFAULT_TEMPLATE_PATH
    if os.path.exists(FALLBACK_GDRIVE_TEMPLATE):
        return FALLBACK_GDRIVE_TEMPLATE
    raise FileNotFoundError("Could not locate AAEON_Template_v2.pptx in resources or Google Drive.")

def find_shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name.strip().lower() == name.strip().lower():
            return sh
    return None

def update_paragraph_text(p, text: str, font_size_pt: float = None, bold: bool = None):
    """Safely updates text of a paragraph, preserving font style when possible."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
        if font_size_pt is not None:
            p.runs[0].font.size = Pt(font_size_pt)
        if bold is not None:
            p.runs[0].font.bold = bold
    else:
        p.text = text
        if font_size_pt is not None:
            p.font.size = Pt(font_size_pt)
        if bold is not None:
            p.font.bold = bold

def set_shape_paragraphs(shape, lines: List[str], default_font_size: float = None):
    """Sets multiple lines for a shape text frame."""
    tf = shape.text_frame
    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
            update_paragraph_text(p, line, font_size_pt=default_font_size)
        else:
            p = tf.add_paragraph()
            update_paragraph_text(p, line, font_size_pt=default_font_size)
    for i in range(len(lines), len(tf.paragraphs)):
        update_paragraph_text(tf.paragraphs[i], "")

def populate_slide_1(slide, data: Dict[str, Any]):
    # Title 1, Subtitle 2
    sh_title = find_shape_by_name(slide, "Title 1") or slide.shapes[0]
    if "title" in data and len(sh_title.text_frame.paragraphs) > 0:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"], font_size_pt=36.0, bold=True)
    if "subtitle" in data and len(sh_title.text_frame.paragraphs) > 1:
        update_paragraph_text(sh_title.text_frame.paragraphs[1], data["subtitle"], font_size_pt=36.0)

    sh_sub = find_shape_by_name(slide, "Subtitle 2")
    if sh_sub and "presenter" in data and sh_sub.text_frame.paragraphs:
        update_paragraph_text(sh_sub.text_frame.paragraphs[0], data["presenter"], font_size_pt=20.0)

def populate_slide_2(slide, data: Dict[str, Any]):
    # Title
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    # 5 Agenda boxes: TextBox 5, 7, 9, 11, 13
    box_names = ["TextBox 5", "TextBox 7", "TextBox 9", "TextBox 11", "TextBox 13"]
    items = data.get("items", [])
    for idx, name in enumerate(box_names):
        if idx < len(items):
            sh = find_shape_by_name(slide, name)
            if sh and len(sh.text_frame.paragraphs) >= 2:
                item = items[idx]
                heading = item.get("heading", f"0{idx+1} 主題")
                desc = item.get("description", "• 請輸入本章節核心摘要與預期效益說明")
                update_paragraph_text(sh.text_frame.paragraphs[0], heading, font_size_pt=13.0, bold=True)
                update_paragraph_text(sh.text_frame.paragraphs[1], desc, font_size_pt=10.5)

def populate_slide_3(slide, data: Dict[str, Any]):
    # Title 1, Subtitle 2
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title:
        ch_num = data.get("chapter_number", "01")
        title = data.get("title", "章節主標題")
        sub_en = data.get("subtitle_en", "Chapter English Title")
        if len(sh_title.text_frame.paragraphs) >= 2:
            update_paragraph_text(sh_title.text_frame.paragraphs[0], f"{ch_num} / {title}", font_size_pt=32.0, bold=True)
            update_paragraph_text(sh_title.text_frame.paragraphs[1], sub_en, font_size_pt=32.0)
    sh_sub = find_shape_by_name(slide, "Subtitle 2")
    if sh_sub and "summary" in data and sh_sub.text_frame.paragraphs:
        update_paragraph_text(sh_sub.text_frame.paragraphs[0], data["summary"], font_size_pt=16.0)

def populate_slide_4(slide, data: Dict[str, Any]):
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    sh_takeaway = find_shape_by_name(slide, "TextBox 4")
    if sh_takeaway and "takeaway" in data:
        if len(sh_takeaway.text_frame.paragraphs) >= 2:
            update_paragraph_text(sh_takeaway.text_frame.paragraphs[0], "【核心結論（Executive Takeaway）】", font_size_pt=14.0, bold=True)
            update_paragraph_text(sh_takeaway.text_frame.paragraphs[1], data["takeaway"], font_size_pt=12.0)

    sh_pillars = find_shape_by_name(slide, "TextBox 6")
    if sh_pillars and "pillars" in data:
        pillars = data.get("pillars", [])
        lines = ["關鍵架構分析重點（Key Architectural Pillars）"]
        for p in pillars:
            lines.append(f"• {p}")
        set_shape_paragraphs(sh_pillars, lines, default_font_size=12.0)
        if sh_pillars.text_frame.paragraphs:
            update_paragraph_text(sh_pillars.text_frame.paragraphs[0], lines[0], font_size_pt=14.0, bold=True)

def populate_slide_5(slide, data: Dict[str, Any]):
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    for opt_key, box_name in [("option_a", "TextBox 4"), ("option_b", "TextBox 6")]:
        if opt_key in data:
            sh = find_shape_by_name(slide, box_name)
            if sh:
                opt = data[opt_key]
                lines = [
                    f"{opt.get('header', '方案')}",
                    f"{opt.get('subtitle', '')}",
                    f"• 【優點（Pros）】{opt.get('pros', '')}",
                    f"• 【缺點（Cons）】{opt.get('cons', '')}",
                    f"• 【BOM 成本】{opt.get('cost', '')}",
                    f"• 【設計規範】{opt.get('spec', '')}",
                    f"• 【建議評估】{opt.get('recommendation', '')}"
                ]
                set_shape_paragraphs(sh, lines, default_font_size=11.0)
                if len(sh.text_frame.paragraphs) > 0:
                    update_paragraph_text(sh.text_frame.paragraphs[0], lines[0], font_size_pt=14.0, bold=True)
                if len(sh.text_frame.paragraphs) > 1:
                    update_paragraph_text(sh.text_frame.paragraphs[1], lines[1], font_size_pt=10.0)

def populate_slide_6(slide, data: Dict[str, Any]):
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    # 3 Pillar boxes: TextBox 4, TextBox 6, TextBox 8
    box_names = ["TextBox 4", "TextBox 6", "TextBox 8"]
    pillars = data.get("pillars", [])
    for idx, name in enumerate(box_names):
        if idx < len(pillars):
            sh = find_shape_by_name(slide, name)
            if sh:
                p_data = pillars[idx]
                lines = [
                    p_data.get("title_zh", f"模組 / 支柱 {idx+1}"),
                    p_data.get("title_en", f"Pillar {idx+1} English")
                ]
                points = p_data.get("points", [])
                for pt in points:
                    lines.append(f" • {pt}")
                set_shape_paragraphs(sh, lines, default_font_size=11.0)
                if len(sh.text_frame.paragraphs) > 0:
                    update_paragraph_text(sh.text_frame.paragraphs[0], lines[0], font_size_pt=14.0, bold=True)
                if len(sh.text_frame.paragraphs) > 1:
                    update_paragraph_text(sh.text_frame.paragraphs[1], lines[1], font_size_pt=10.0)

def populate_slide_7(slide, data: Dict[str, Any]):
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    # 4 Risk boxes: TextBox 4, TextBox 6, TextBox 8, TextBox 10
    box_names = ["TextBox 4", "TextBox 6", "TextBox 8", "TextBox 10"]
    default_headers = [
        "【高衝擊 / 高機率】關鍵技術風險",
        "【高衝擊 / 低機率】平台相容風險",
        "【低衝擊 / 高機率】機構公差風險",
        "【低衝擊 / 低機率】常規物料交期"
    ]
    risks = data.get("risks", [])
    for idx, name in enumerate(box_names):
        if idx < len(risks):
            sh = find_shape_by_name(slide, name)
            if sh:
                r = risks[idx]
                header = r.get("header", default_headers[idx])
                risk_desc = r.get("risk", "風險描述")
                mitigation = r.get("mitigation", "因應對策")
                lines = [
                    header,
                    f"風險點：{risk_desc}",
                    f"因應對策：{mitigation}"
                ]
                set_shape_paragraphs(sh, lines, default_font_size=10.5)
                if len(sh.text_frame.paragraphs) > 0:
                    update_paragraph_text(sh.text_frame.paragraphs[0], header, font_size_pt=12.0, bold=True)

def populate_slide_8(slide, data: Dict[str, Any]):
    sh_title = find_shape_by_name(slide, "Title 1")
    if sh_title and "title" in data:
        update_paragraph_text(sh_title.text_frame.paragraphs[0], data["title"])

    # 4 Phase boxes: TextBox 4, TextBox 6, TextBox 8, TextBox 10
    box_names = ["TextBox 4", "TextBox 6", "TextBox 8", "TextBox 10"]
    phases = data.get("phases", [])
    for idx, name in enumerate(box_names):
        if idx < len(phases):
            sh = find_shape_by_name(slide, name)
            if sh:
                ph = phases[idx]
                phase_num = ph.get("phase", f"Phase {idx+1}：Q{idx+1} 2026")
                title_zh = ph.get("title_zh", "階段任務標題")
                title_en = ph.get("title_en", "Phase English Title")
                lines = [phase_num, title_zh, title_en, ""]
                tasks = ph.get("tasks", [])
                for t in tasks:
                    lines.append(f"• {t}")
                set_shape_paragraphs(sh, lines, default_font_size=10.0)
                if len(sh.text_frame.paragraphs) > 0:
                    update_paragraph_text(sh.text_frame.paragraphs[0], phase_num, font_size_pt=13.0, bold=True)
                if len(sh.text_frame.paragraphs) > 1:
                    update_paragraph_text(sh.text_frame.paragraphs[1], title_zh, font_size_pt=12.0, bold=True)
                if len(sh.text_frame.paragraphs) > 2:
                    update_paragraph_text(sh.text_frame.paragraphs[2], title_en, font_size_pt=9.0)

def populate_slide_9(slide, data: Dict[str, Any]):
    # Text Placeholder 1 or first text shape
    sh = find_shape_by_name(slide, "Text Placeholder 1") or slide.shapes[0]
    if sh:
        lines = [
            data.get("main_text", "Thank You"),
            data.get("sub_text", "Q&A 與討論交流"),
            data.get("company", "研揚科技 AAEON Technology Inc.  |  研發中心")
        ]
        set_shape_paragraphs(sh, lines, default_font_size=18.0)
        if len(sh.text_frame.paragraphs) > 0:
            update_paragraph_text(sh.text_frame.paragraphs[0], lines[0], font_size_pt=40.0, bold=True)
        if len(sh.text_frame.paragraphs) > 1:
            update_paragraph_text(sh.text_frame.paragraphs[1], lines[1], font_size_pt=22.0)
        if len(sh.text_frame.paragraphs) > 2:
            update_paragraph_text(sh.text_frame.paragraphs[2], lines[2], font_size_pt=14.0)

POPULATORS = {
    1: populate_slide_1,
    2: populate_slide_2,
    3: populate_slide_3,
    4: populate_slide_4,
    5: populate_slide_5,
    6: populate_slide_6,
    7: populate_slide_7,
    8: populate_slide_8,
    9: populate_slide_9
}

def generate(payload_path: str, output_path: str, template_path: str = None):
    with open(payload_path, "r", encoding="utf-8") as f:
        payload = json.load(f)

    tpl_path = get_template_path(template_path)
    print(f"Loading template from: {tpl_path}")
    prs = pptx.Presentation(tpl_path)

    slides_payload = payload.get("slides", {})
    if isinstance(slides_payload, list):
        slides_dict = {str(i+1): s for i, s in enumerate(slides_payload)}
    else:
        slides_dict = {str(k): v for k, v in slides_payload.items()}

    for slide_num, populator in POPULATORS.items():
        slide_key = str(slide_num)
        if slide_key in slides_dict and slide_num <= len(prs.slides):
            print(f"Populating Slide {slide_num}...")
            populator(prs.slides[slide_num - 1], slides_dict[slide_key])

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)
    print(f"Successfully generated presentation: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Generate AAEON Presentation from JSON payload.")
    parser.add_argument("--input", "-i", required=True, help="Path to JSON payload file")
    parser.add_argument("--output", "-o", required=True, help="Output PPTX path")
    parser.add_argument("--template", "-t", default=None, help="Custom PPTX template path (optional)")

    args = parser.parse_args()
    generate(args.input, args.output, args.template)

if __name__ == "__main__":
    main()
